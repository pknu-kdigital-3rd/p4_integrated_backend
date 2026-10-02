"""Build the Korean Project 4 deliverables from local project evidence.

Run with python-docx, matplotlib and Pillow installed. The input model archive
can be overridden with P4_MODEL_ARCHIVE. Assets and metrics are saved beside
this script so figures and editable tables can be revised together.
"""
from __future__ import annotations

import csv
import json
import os
import shutil
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parents[1]
ASSETS = Path(__file__).resolve().parent / "assets"
DATA = Path(__file__).resolve().parent / "data"
ARCHIVE = Path(os.environ.get("P4_MODEL_ARCHIVE", "/dat/3rdgen_archive/project4/yolo_carafe_aspp"))
RUNS = ARCHIVE / "runs/segment/experiments/results"
PROJECT = "AI 기반 지능형 교통·차량 통합 관제 시스템"
NAVY = "233D60"
TEAL = "08758F"
TEXT = "202C39"
FONT = "맑은 고딕"
COLORS = ["#42566F", "#16837C", "#D99527", "#925BA6", "#086B9A"]
LABELS = ["A0 기본 모델", "A1 CARAFE", "A2 ASPP", "A3 CARAFE + ASPP", "A4 CARAFE + ASPP + Decoder"]
RUN_NAMES = ["A0_local", "A1_local", "A2_local", "A3_local", "A4_local-2"]
METRICS = ["metrics/precision(B)", "metrics/recall(B)", "metrics/mAP50(B)", "metrics/mAP50-95(B)",
           "metrics/precision(M)", "metrics/recall(M)", "metrics/mAP50(M)", "metrics/mAP50-95(M)"]
ALL_ROWS: list[list[dict]] = []
BEST: list[dict] = []
ASSETS.mkdir(parents=True, exist_ok=True)
DATA.mkdir(parents=True, exist_ok=True)
font_path = "/Applications/Microsoft Word.app/Contents/Resources/DFonts/malgun.ttf"
if not Path(font_path).exists():
    font_path = "/Library/Fonts/Arial Unicode.ttf"
font_manager.fontManager.addfont(font_path)
plt.rcParams.update({"font.family": font_manager.FontProperties(fname=font_path).get_name(),
                     "axes.unicode_minus": False, "font.size": 10, "savefig.dpi": 220})


def prepare_metrics():
    for run in RUN_NAMES:
        path = RUNS / run / "results.csv"
        with path.open() as f:
            rows = [{k.strip(): float(v.strip()) for k, v in row.items()} for row in csv.DictReader(f)]
        ALL_ROWS.append(rows)
        BEST.append(max(rows, key=lambda r: r["metrics/mAP50-95(M)"]))
        shutil.copyfile(path, DATA / f"{run}_results.csv")
    payload = {run: {"best_mask_epoch": int(best["epoch"]), "recorded_epochs": len(rows),
                     "metrics": {k: best[k] for k in METRICS}}
               for run, rows, best in zip(RUN_NAMES, ALL_ROWS, BEST)}
    (DATA / "segmentation_metrics.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2))


def plot_training():
    for idx, (rows, best) in enumerate(zip(ALL_ROWS, BEST)):
        fig, axes = plt.subplots(2, 2, figsize=(9.2, 6.8), constrained_layout=True)
        epoch = [r["epoch"] for r in rows]
        specs = [
            ("분할 정확도·재현율", ["metrics/precision(M)", "metrics/recall(M)"], ["Precision", "Recall"], "값"),
            ("분할 mAP", ["metrics/mAP50(M)", "metrics/mAP50-95(M)"], ["mAP50", "mAP50-95"], "값"),
            ("분할 손실", ["train/seg_loss", "val/seg_loss"], ["학습", "검증"], "Loss"),
            ("박스 mAP", ["metrics/mAP50(B)", "metrics/mAP50-95(B)"], ["mAP50", "mAP50-95"], "값"),
        ]
        for ax, (title, keys, legends, ylabel) in zip(axes.flat, specs):
            for k, legend, color in zip(keys, legends, ["#08758F", "#C77825"]):
                ax.plot(epoch, [r[k] for r in rows], lw=1.6, label=legend, color=color)
            ax.axvline(best["epoch"], color="#8593A1", lw=1, ls="--", label=f"최고 Mask epoch {int(best['epoch'])}")
            ax.set(title=title, xlabel="Epoch", ylabel=ylabel, xlim=(1, 100))
            ax.grid(alpha=.18)
            ax.legend(fontsize=8, loc="best", frameon=False)
            ax.spines[["top", "right"]].set_visible(False)
        fig.suptitle(f"{LABELS[idx]} · 저장된 학습 기록", fontsize=14, color="#233D60")
        fig.savefig(ASSETS / f"training_A{idx}.png", facecolor="white")
        plt.close(fig)
    fig, ax = plt.subplots(figsize=(9, 4.1), constrained_layout=True)
    markers = ["o", "s", "^", "D", "v"]
    for i, rows in enumerate(ALL_ROWS):
        ax.plot([r["epoch"] for r in rows], [r["metrics/mAP50-95(M)"] for r in rows],
                label=f"A{i}", color=COLORS[i], lw=1.5, marker=markers[i], ms=4,
                markevery=(i * 2, 17), ls=["-", "-", "--", ":", "-"][i])
    ax.set(xlabel="Epoch", ylabel="Mask mAP50-95", xlim=(1, 100), title="A0–A4 분할 성능 변화")
    ax.grid(alpha=.18)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(ncol=5, frameon=False)
    fig.savefig(ASSETS / "training_comparison.png", facecolor="white")
    plt.close(fig)
    fig, ax = plt.subplots(figsize=(8.5, 3.6), constrained_layout=True)
    vals = [r["metrics/mAP50-95(M)"] for r in BEST]
    bars = ax.bar([f"A{i}" for i in range(5)], vals, color=COLORS, width=.58)
    for bar, val in zip(bars, vals):
        ax.text(bar.get_x() + bar.get_width()/2, val+.003, f"{val:.5f}", ha="center", fontsize=11)
    ax.set(ylim=(0, .46), ylabel="Mask mAP50-95", title="실험별 최고 Mask mAP50-95")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", alpha=.15)
    ax.set_axisbelow(True)
    fig.savefig(ASSETS / "best_comparison.png", facecolor="white")
    plt.close(fig)


def box(ax, xy, w, h, text, color="#EDF4F7", size=10):
    ax.add_patch(FancyBboxPatch(xy, w, h, boxstyle="round,pad=0.08,rounding_size=0.10", fc=color, ec="#9EB2C2", lw=1))
    ax.text(xy[0]+w/2, xy[1]+h/2, text, ha="center", va="center", fontsize=size, color="#233D60")


def arrow(ax, start, end, label="", pos=None):
    ax.annotate("", xy=end, xytext=start, arrowprops={"arrowstyle": "->", "color": "#637F96", "lw": 1.3})
    if label:
        x, y = pos or ((start[0]+end[0])/2, (start[1]+end[1])/2+.12)
        ax.text(x, y, label, ha="center", va="center", fontsize=8.5, color="#4D6375",
                bbox={"facecolor": "white", "edgecolor": "none", "pad": 1.5})


def diagram_arch():
    fig, ax = plt.subplots(figsize=(10, 6.3))
    ax.set(xlim=(0, 10), ylim=(0, 7));ax.axis("off")
    box(ax, (.3, 5.6), 2.5, .8, "관제 웹\nLeaflet · Live View")
    box(ax, (7.2, 5.6), 2.5, .8, "Android 단말\n카메라 · GPS/IMU")
    box(ax, (3.5, 4.6), 3, .8, "Nginx / Coturn\nHTTPS · TURN", "#DDEDF2")
    box(ax, (.3, 3), 2.5, 1, "Node / Express\n인증 · 업무 · 가상 배차")
    box(ax, (3.8, 3), 2.4, 1, "Routing / Tracking\nBIMS · A* · 도로 제한")
    box(ax, (7.2, 3.6), 2.5, .7, "Go Media Relay\n영상 수신 · 녹화")
    box(ax, (7.2, 2.1), 2.5, .8, "Vision / FastAPI\nYOLO26-seg · UniDepth")
    box(ax, (.3, .8), 2.5, .8, "PostgreSQL / PostGIS\n업무 · 위치 · 메타데이터", "#E9EDEB")
    box(ax, (3.8, .8), 2.4, .8, "MinIO\n녹화 영상 객체", "#E9EDEB")
    box(ax, (3.8, 6.0), 2.4, .55, "BIMS / OSM / 제한 데이터", "#F7F1E8", size=8.5)
    arrow(ax, (2.8, 6), (3.5, 5.1), "HTTPS", (3.1,5.75))
    arrow(ax, (7.2, 6), (6.5, 5.1), "HTTPS / WebRTC", (7.05,5.55))
    arrow(ax, (3.8, 4.6), (1.8, 4), "업무 API", (2.55,4.6))
    arrow(ax, (6.5, 4.8), (8.45, 4.3), "미디어", (7.4,4.65))
    arrow(ax, (2.8, 3.5), (3.8, 3.5), "HTTP")
    arrow(ax, (5, 6), (5, 4), "수집·조회", (5.5,5.7))
    arrow(ax, (8.45, 3.6), (8.45, 2.9), "Unix 소켓 / HTTP", (8.5,3.24))
    arrow(ax, (1.55, 3), (1.55, 1.6), "영속 저장", (1.55,2.3))
    arrow(ax, (7.2, 3.8), (6.2, 1.6), "S3 업로드", (6.4,2.4))
    arrow(ax, (7.2, 2.3), (2.8, 3.1), "탐지 샘플 → Node 내부 API", (4.9,2.7))
    ax.text(5, .22, "영상·추론과 업무 처리를 서비스 경계에서 연결", ha="center", fontsize=11, color="#08758F")
    fig.tight_layout();fig.savefig(ASSETS / "system_architecture.png", facecolor="white");plt.close(fig)


def diagram_model():
    fig, ax = plt.subplots(figsize=(10, 5.8))
    ax.set(xlim=(0, 10), ylim=(0, 6));ax.axis("off")
    box(ax, (.2, 2.7), 1.5, .85, "입력 영상\nRGB", size=10)
    box(ax, (2.1, 2.4), 1.7, 1.45, "YOLO26s\nBackbone\nP2 / P3 / P4 / P5")
    box(ax, (4.35, 3.65), 2.0, 1.0, "P5 ASPP\nrate = 3, 6, 9", "#DDEDF2")
    box(ax, (6.95, 3.65), 2.65, 1.0, "CARAFE Neck\nP5 → P4 → P3", "#DDEDF2")
    box(ax, (6.95, 1.8), 2.65, 1.0, "P3 / P4 / P5\n박스·클래스·마스크 계수")
    box(ax, (4.35, .3), 2.1, 1.25, "P2 상세 특징 + P3\nCARAFE · 특징 결합\n3×3 정제 · Proto", "#E9EDEB", size=9.5)
    box(ax, (7.4, .3), 2.2, .9, "Instance Mask\n객체별 분할 결과", "#E9EDEB")
    arrow(ax, (1.7, 3.15), (2.1, 3.15))
    arrow(ax, (3.8, 3.4), (4.35, 4.1), "P5", (4.08,3.93))
    arrow(ax, (6.35, 4.15), (6.95, 4.15))
    arrow(ax, (8.25, 3.65), (8.25, 2.8))
    arrow(ax, (3.0, 2.4), (4.35, .93), "P2 / stride 4", (3.2,1.45))
    arrow(ax, (6.95, 3.85), (5.4, 1.55), "P3 / stride 8", (5.8,2.65))
    arrow(ax, (6.45, .75), (7.4, .75), "32 Proto", (6.9,1.14))
    arrow(ax, (8.5, 1.8), (8.5, 1.2))
    ax.text(5, 5.45, "Custom YOLO26s-seg A4 구조", ha="center", fontsize=15, color="#233D60")
    fig.tight_layout();fig.savefig(ASSETS / "custom_yolo_architecture.png", facecolor="white");plt.close(fig)


def diagram_domain():
    fig, ax = plt.subplots(figsize=(10, 5.8))
    ax.set(xlim=(0, 10), ylim=(0, 6));ax.axis("off")
    box(ax, (3.75, 4.9), 2.5, .65, "vehicle · 공통 차량", "#DDEDF2")
    box(ax, (.35, 3.65), 4.05, .7, "실제 운행 · trip / route")
    box(ax, (5.6, 3.65), 4.05, .7, "가상 운행 · virtual_trip / virtual_route")
    box(ax, (.35, 2.15), 4.05, .9, "vehicle_position\n실제 GPS · 출처 · 원본 시각")
    box(ax, (5.6, 2.15), 4.05, .9, "virtual_vehicle_state\n현재 간선 · 진행 거리 · 시뮬레이션 상태")
    box(ax, (.35, .5), 4.05, 1.0, "trip_video / trip_video_detection_sample\n영상 PTS · epoch · frame_seq")
    box(ax, (5.6, .5), 4.05, 1.0, "virtual_scenario / dispatch_request\nwaypoint · restriction · operator_event")
    arrow(ax, (4.0, 4.9), (2.4, 4.35));arrow(ax, (6.0, 4.9), (7.6, 4.35))
    arrow(ax, (2.35, 3.65), (2.35, 3.05));arrow(ax, (7.6, 3.65), (7.6, 3.05))
    arrow(ax, (2.35, 2.15), (2.35, 1.5));arrow(ax, (7.6, 2.15), (7.6, 1.5))
    fig.tight_layout();fig.savefig(ASSETS / "domain_structure.png", facecolor="white");plt.close(fig)


def plot_classes():
    names = ["person", "bicycle", "car", "motorcycle", "bus", "truck"]
    values = [[.466,.201,.375,.368,.650,.348], [.465,.196,.379,.378,.649,.356], [.468,.212,.382,.388,.654,.356]]
    fig, ax = plt.subplots(figsize=(9, 4.1), constrained_layout=True)
    x = list(range(len(names)))
    for off, vals, label, color in zip([-.25,0,.25], values, ["A0", "A4", "공식 모델"], [COLORS[0], COLORS[4], "#67A69A"]):
        ax.bar([i+off for i in x], vals, width=.23, label=label, color=color)
    ax.set(xticks=x, xticklabels=names, ylabel="Mask AP50-95", ylim=(0,.75), title="주요 교통 객체 클래스별 분할 성능")
    ax.legend(frameon=False, ncol=3);ax.grid(axis="y", alpha=.15);ax.set_axisbelow(True)
    ax.spines[["top","right"]].set_visible(False)
    fig.savefig(ASSETS / "class_comparison.png", facecolor="white");plt.close(fig)


def shade(cell, fill):
    el = OxmlElement("w:shd");el.set(qn("w:fill"), fill);cell._tc.get_or_add_tcPr().append(el)


def field(paragraph, instr, cached="1"):
    r = paragraph.add_run()
    b = OxmlElement("w:fldChar");b.set(qn("w:fldCharType"), "begin");r._r.append(b)
    r = paragraph.add_run();t = OxmlElement("w:instrText");t.set(qn("xml:space"), "preserve");t.text = f" {instr} ";r._r.append(t)
    r = paragraph.add_run();s = OxmlElement("w:fldChar");s.set(qn("w:fldCharType"), "separate");r._r.append(s)
    paragraph.add_run(cached)
    r = paragraph.add_run();e = OxmlElement("w:fldChar");e.set(qn("w:fldCharType"), "end");r._r.append(e)


class Writer:
    def __init__(self, kind):
        self.kind = kind
        self.doc = Document()
        self.page_num = 0
        self.chapters = []
        self.tables = 0
        self.figures = 0
        self.bookmark_id = 1
        self.toc_holder = None
        styles = self.doc.styles
        for name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "Caption", "Header", "Footer"]:
            st = styles[name];st.font.name=FONT;st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),FONT)
            st.font.color.rgb=RGBColor.from_string(TEXT)
        norm=styles["Normal"];norm.font.size=Pt(10.5)
        norm.paragraph_format.line_spacing=1.28
        norm.paragraph_format.space_after=Pt(7)
        norm.paragraph_format.widow_control=True
        for name, size, color in [("Heading 1",20.5,NAVY),("Heading 2",14.2,TEAL),("Heading 3",11.3,TEAL)]:
            st=styles[name];st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(color)
            st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(8)
            st.paragraph_format.keep_with_next=True
        styles["Caption"].font.size=Pt(8.5);styles["Caption"].font.color.rgb=RGBColor.from_string(NAVY)
        styles["Caption"].paragraph_format.space_after=Pt(5)
        self.doc.core_properties.title=f"4차 프로젝트 {kind} | {PROJECT}"
        self.doc.core_properties.subject="ITS 통합 관제 시스템 · Custom YOLO26-seg · UniDepth"
        self.doc.core_properties.author=""
        self.doc.core_properties.keywords="ITS, 차량 관제, YOLO26-seg, CARAFE, ASPP, UniDepth"
        self.setup_section(self.doc.sections[0])
        settings=self.doc.settings.element
        lang=OxmlElement("w:themeFontLang");lang.set(qn("w:val"),"ko-KR");lang.set(qn("w:eastAsia"),"ko-KR");settings.append(lang)
        compat=OxmlElement("w:doNotAutoHyphenate");settings.append(compat)

    def setup_section(self, sec):
        sec.page_width=Cm(21);sec.page_height=Cm(29.7)
        sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7)
        sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
        sec.header_distance=Cm(.75);sec.footer_distance=Cm(.8)

    def cover(self):
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(85)
        p=self.doc.add_paragraph("4차 프로젝트")
        p.paragraph_format.space_after=Pt(10);r=p.runs[0];r.font.size=Pt(19);r.font.bold=True;r.font.color.rgb=RGBColor.from_string(TEAL)
        p=self.doc.add_paragraph("기획서" if self.kind=="기획서" else "결과 보고서")
        p.paragraph_format.space_after=Pt(24);r=p.runs[0];r.font.size=Pt(32);r.font.bold=True;r.font.color.rgb=RGBColor.from_string(TEAL)
        p=self.doc.add_paragraph("AI 기반 지능형 교통·차량\n통합 관제 시스템")
        p.paragraph_format.line_spacing=1.4;r=p.runs[0];r.font.size=Pt(22);r.font.color.rgb=RGBColor.from_string(NAVY)
        p=self.doc.add_paragraph("차량 관제 · 경로 탐색 · 영상 분석 · 운행 기록")
        p.paragraph_format.space_after=Pt(90);r=p.runs[0];r.font.size=Pt(11);r.font.color.rgb=RGBColor.from_string(TEAL)
        t=self.doc.add_table(rows=3,cols=2);t.alignment=WD_TABLE_ALIGNMENT.RIGHT;t.autofit=False
        t.columns[0].width=Cm(2.1);t.columns[1].width=Cm(7)
        for row,label in zip(t.rows,["팀원","개발 기간","제출일"]):
            row.cells[0].text=label;row.cells[1].text=""
            row.height=Cm(.7);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
            for c in row.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.space_after=Pt(0)
                    for rr in pp.runs:rr.font.size=Pt(10);rr.font.color.rgb=RGBColor.from_string(TEAL)
            borders=OxmlElement("w:tcBorders");bottom=OxmlElement("w:bottom");bottom.set(qn("w:val"),"single");bottom.set(qn("w:sz"),"4");bottom.set(qn("w:color"),"BDCED8");borders.append(bottom);row.cells[1]._tc.get_or_add_tcPr().append(borders)
        sec=self.doc.add_section(WD_SECTION_START.NEW_PAGE);self.setup_section(sec)
        sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
        h=sec.header.paragraphs[0];h.text=f"4차 프로젝트 {self.kind} | {PROJECT}"
        h.paragraph_format.space_after=Pt(0)
        for r in h.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string(TEAL)
        f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        f.add_run("페이지 ");field(f,"PAGE");f.add_run(" / ");field(f,"SECTIONPAGES","1")
        for r in f.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string("718394")
        pg=OxmlElement("w:pgNumType");pg.set(qn("w:start"),"1");sec._sectPr.append(pg)

    def page(self, heading, major=False):
        break_before = self.page_num > 0
        self.page_num += 1
        p=self.doc.add_paragraph(heading,"Heading 1" if major else "Heading 2")
        p.paragraph_format.page_break_before = break_before
        if major:
            self.chapters.append((heading,self.page_num,f"chapter_{self.bookmark_id}"))
            b=OxmlElement("w:bookmarkStart");b.set(qn("w:id"),str(self.bookmark_id));b.set(qn("w:name"),f"chapter_{self.bookmark_id}");p._p.insert(0,b)
            e=OxmlElement("w:bookmarkEnd");e.set(qn("w:id"),str(self.bookmark_id));p._p.append(e);self.bookmark_id+=1

    def h(self, text):self.doc.add_paragraph(text,"Heading 2")
    def h3(self, text):self.doc.add_paragraph(text,"Heading 3")
    def p(self, text):self.doc.add_paragraph(text)
    def bullet(self, text):
        p=self.doc.add_paragraph("•  "+text);p.paragraph_format.left_indent=Cm(.25);p.paragraph_format.first_line_indent=Cm(-.2)
        p.paragraph_format.space_after=Pt(5)
    def steps(self, rows):
        for i,text in enumerate(rows,1):
            p=self.doc.add_paragraph(f"{i}. {text}");p.paragraph_format.left_indent=Cm(.4);p.paragraph_format.first_line_indent=Cm(-.35)
            p.paragraph_format.space_after=Pt(6)
    def note(self, text):
        p=self.doc.add_paragraph(text);p.paragraph_format.space_after=Pt(7)
        for r in p.runs:r.font.size=Pt(8.2);r.font.color.rgb=RGBColor.from_string("627585")

    def table(self, title, headers, rows, widths=None, blank=False, font_size=9):
        self.tables+=1
        self.doc.add_paragraph(f"표 {self.tables}. {title}","Caption")
        t=self.doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        widths=widths or [17.4/len(headers)]*len(headers)
        for c,w in zip(t.columns,widths):c.width=Cm(w)
        for c,head,w in zip(t.rows[0].cells,headers,widths):c.text=str(head);c.width=Cm(w);shade(c,NAVY)
        repeat=OxmlElement("w:tblHeader");t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for i,values in enumerate(rows):
            row=t.add_row()
            for c,v,w in zip(row.cells,values,widths):
                c.text=str(v);c.width=Cm(w)
                shade(c,"F0F4F7" if i%2==0 else "FFFFFF")
            if blank:
                row.height=Cm(.88);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
            cant=OxmlElement("w:cantSplit");row._tr.get_or_add_trPr().append(cant)
        for ri,row in enumerate(t.rows):
            for c in row.cells:
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=c._tc.get_or_add_tcPr();margins=OxmlElement("w:tcMar")
                for side,value in [("top",70),("bottom",70),("left",85),("right",85)]:
                    e=OxmlElement("w:"+side);e.set(qn("w:w"),str(value));e.set(qn("w:type"),"dxa");margins.append(e)
                tcpr.append(margins)
                borders=OxmlElement("w:tcBorders")
                for side in ["top","bottom","left","right"]:
                    e=OxmlElement("w:"+side);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"CDD8E1");borders.append(e)
                tcpr.append(borders)
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15
                    for r in p.runs:
                        r.font.name=FONT;r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),FONT)
                        r.font.size=Pt(font_size);r.font.bold=ri==0
                        r.font.color.rgb=RGBColor.from_string("FFFFFF" if ri==0 else TEXT)
        self.doc.add_paragraph().paragraph_format.space_after=Pt(0)
        return t

    def image(self, name, title, width=17.0):
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(4)
        shape=p.add_run().add_picture(str(ASSETS/name),width=Cm(width))
        shape._inline.docPr.set("descr",title)
        self.figures+=1
        self.doc.add_paragraph(f"그림 {self.figures}. {title}","Caption")

    def blank_area(self, title, height=3.8):
        self.doc.add_paragraph(title,"Caption")
        t=self.doc.add_table(rows=1,cols=1);t.columns[0].width=Cm(17.4)
        t.rows[0].height=Cm(height);t.rows[0].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        shade(t.cell(0,0),"F8FAFB")
        borders=OxmlElement("w:tcBorders")
        for side in ["top","bottom","left","right"]:
            b=OxmlElement("w:"+side);b.set(qn("w:val"),"single");b.set(qn("w:sz"),"4");b.set(qn("w:color"),"CCD9E2");borders.append(b)
        t.cell(0,0)._tc.get_or_add_tcPr().append(borders)

    def toc(self):
        self.page("결과 보고서 목차")
        p=self.doc.add_paragraph("PROJECT REPORT");p.runs[0].font.color.rgb=RGBColor.from_string(TEAL)
        self.toc_holder=self.doc.add_paragraph()
        self.p("본 보고서는 통합 시스템의 구성과 구현 결과, Custom YOLO26-seg의 학습·평가 자료를 정리한다.")

    def fill_toc(self):
        if self.toc_holder is None:return
        holder=self.toc_holder
        r=OxmlElement("w:r");b=OxmlElement("w:fldChar");b.set(qn("w:fldCharType"),"begin");r.append(b);holder._p.append(r)
        r=OxmlElement("w:r");instr=OxmlElement("w:instrText");instr.set(qn("xml:space"),"preserve");instr.text=' TOC \\o "1-1" \\h \\z \\u ';r.append(instr);holder._p.append(r)
        r=OxmlElement("w:r");b=OxmlElement("w:fldChar");b.set(qn("w:fldCharType"),"separate");r.append(b);holder._p.append(r)
        anchor=holder._p
        for title,page,bookmark in self.chapters:
            p=OxmlElement("w:p");ppr=OxmlElement("w:pPr")
            sp=OxmlElement("w:spacing");sp.set(qn("w:after"),"220");ppr.append(sp)
            tabs=OxmlElement("w:tabs");tab=OxmlElement("w:tab");tab.set(qn("w:val"),"right");tab.set(qn("w:leader"),"dot");tab.set(qn("w:pos"),"9400");tabs.append(tab);ppr.append(tabs);p.append(ppr)
            link=OxmlElement("w:hyperlink");link.set(qn("w:anchor"),bookmark)
            r=OxmlElement("w:r");rpr=OxmlElement("w:rPr");color=OxmlElement("w:color");color.set(qn("w:val"),NAVY);rpr.append(color);r.append(rpr)
            t=OxmlElement("w:t");t.text=title;r.append(t);link.append(r);p.append(link)
            r=OxmlElement("w:r");r.append(OxmlElement("w:tab"));t=OxmlElement("w:t");t.text=str(page);r.append(t);p.append(r)
            anchor.addnext(p);anchor=p
        p=OxmlElement("w:p");r=OxmlElement("w:r");b=OxmlElement("w:fldChar");b.set(qn("w:fldCharType"),"end");r.append(b);p.append(r);anchor.addnext(p)

    def save(self, name):
        self.fill_toc()
        self.doc.save(OUT/name)
        return {"file":name,"planned_body_pages":self.page_num,"tables":self.tables,"figures":self.figures,
                "chapters":[{"title":t,"page":p} for t,p,_ in self.chapters]}


STAKEHOLDERS = [
    ["관리자", "계정·차량 관리, 권한·설정 관리", "데이터 정합성과 접근 통제"],
    ["관제 운영자", "차량·경로·영상 관제, 운행·가상 배차", "상황 파악과 대응"],
    ["조회자", "지도·영상·운행 기록 조회", "읽기 전용 정보 확인"],
    ["Android 단말", "배정 운행 확인, 카메라·GPS/IMU 송출", "운행과 송출의 안정성"],
]
FEATURES = [
    ["인증·권한", "로그인, 역할별 API 접근", "Node/Express"],
    ["차량·운행", "차량 등록, 출발·도착지 설정, 운행 시작·종료", "Node + PostgreSQL/PostGIS"],
    ["위치·경로", "BIMS/단말 위치, 차량 제약을 반영한 A*", "Routing/Tracking"],
    ["영상·AI", "Android 영상, 객체 박스·분할·거리 표시", "Media Relay + Vision"],
    ["녹화·재생", "영상 세그먼트, 탐지·GPS 시간 동기화", "Relay + Node + MinIO + Vision"],
    ["가상 배차", "경로 미리보기, 배차 수락, 통제·우회", "Node + Routing"],
    ["운영", "HTTPS, TURN, 서비스별 실행 환경", "Nginx + Coturn + Compose"],
]
LAYERS = [
    ["클라이언트", "관제 웹 / Android", "지도·영상 표시 / 카메라·센서 송출"],
    ["외부 접점", "Nginx / Coturn", "TLS 종단, 경로 라우팅, NAT 통과"],
    ["업무 처리", "Node / Express", "인증·권한, 업무 API, 가상 시뮬레이션"],
    ["계산", "Routing / Tracking", "도로 그래프, A*, BIMS·단말 위치 병합"],
    ["미디어·비전", "Go Relay / FastAPI Vision", "영상 수신·녹화, 객체 분할·거리 추정"],
    ["저장", "PostgreSQL / PostGIS / MinIO", "업무·공간 데이터, 영상 객체"],
]
STACK = [
    ["관제 웹", "JavaScript, HTML/CSS, Leaflet", "지도·워크스페이스·영상 연동"],
    ["업무 백엔드", "TypeScript, Express 5, Prisma 7", "API, 인증, 업무 데이터 관리"],
    ["경로·추적", "Python, FastAPI, OSM 도로 그래프", "차량 제한 판정·A*·위치 병합"],
    ["미디어 릴레이", "Go 1.22, Pion WebRTC v4", "H.264 수신·전달·녹화"],
    ["비전", "Python 3.12, PyTorch, Ultralytics", "분할·추적·거리 추정"],
    ["차량 단말", "Kotlin, CameraX, WebRTC, OkHttp", "카메라·GPS/IMU·운행 제어"],
    ["업무 저장소", "PostgreSQL 17, PostGIS 3.5", "업무·공간 데이터"],
    ["객체 저장소", "MinIO S3 API", "녹화 영상 세그먼트"],
    ["인프라", "Docker Compose, Nginx, Coturn", "컨테이너 실행·HTTPS·TURN"],
]
CLASSES = [["0","person","사람"],["1","bicycle","자전거"],["2","car","승용차"],
           ["3","motorcycle","오토바이"],["5","bus","버스"],["7","truck","트럭"]]
BLANK_MODEL_ROWS = [[v,""] for v in ["모델명 / 버전","모델 구조","입력 / 출력","데이터셋","학습·검증 분할","입력 해상도","학습 epoch","배치 크기","Optimizer / 학습률","가중치 / 실행 환경"]]
BLANK_EVAL_ROWS = [[v,"",""] for v in ["평가 데이터셋","정답 거리 취득 방식","평가 거리 범위","평가 장비 / 해상도","거리 오차 지표","정확도 지표","추론 시간","장면별 결과"]]


def planning():
    w=Writer("기획서");w.cover()
    w.page("1. 프로젝트 개요",True)
    w.h("1.1 프로젝트명");w.p(PROJECT)
    w.h("1.2 프로젝트 소개")
    w.p("부산 지역의 차량 위치, 계획 경로, 차량 카메라 영상과 운행 기록을 하나의 웹 관제 화면에서 확인하는 시스템을 개발한다. 차량 제원을 반영한 경로 탐색과 가상 배차·도로 통제 시뮬레이션을 연결하고, 영상에서는 Custom YOLO26-seg로 객체를 분할하며 UniDepth로 거리를 측정한다.")
    w.h("1.3 개발 배경")
    w.p("차량 운영에서는 지도상의 위치만으로 도로 상황을 충분히 파악하기 어렵다. 카메라 영상, 경로와 이동 궤적, 과거 운행 기록을 함께 확인할 수 있어야 한다. 특히 대형 차량은 높이·중량·너비와 회전 제한을 반영한 경로가 필요하며, 도로 정체나 통제에 따른 대응을 실제 운행과 분리된 환경에서 검토할 수 있어야 한다.")
    w.h("1.4 프로젝트 목표")
    for t in ["여러 출처의 차량 위치와 경로를 통합 지도에서 제공한다.","Android 영상과 객체 분할·거리 정보를 관제 화면에 연결한다.","녹화 영상·탐지 결과·GPS를 같은 시간축으로 재생한다.","가상 차량의 배차, 경유지 변경과 통제 상황의 우회를 시연한다."]:
        w.bullet(t)
    w.table("주요 사용자와 역할",["사용자","역할","관심사"],STAKEHOLDERS,[3,8,6.4])

    w.page("2. 요구사항 분석",True)
    w.h("2.1 기능 요구사항")
    w.table("기능 영역별 요구사항",["영역","주요 기능","담당 서비스"],FEATURES,[2.4,8.3,6.7])
    w.h("2.2 비기능 요구사항")
    for t in ["업무 데이터는 Node가 관리하고, 경로 계산·일시적 위치 추적은 Routing이 담당한다.","영상·추론 처리는 업무 API와 분리하고 큐를 제한하여 지연 누적을 줄인다.","사용자 역할과 내부 서비스 인증을 적용하고 외부 통신은 HTTPS로 제공한다.","가상 운행은 실제 GPS·영상·녹화 데이터와 분리하며, 서버에서 상태를 관리한다."]:
        w.bullet(t)
    w.h("2.3 사용자 요구사항")
    w.p("운영자는 차량 선택, 경로 확인, Live View, 운행 기록 조회와 가상 배차를 한 화면에서 수행할 수 있어야 한다. 조회자는 읽기 전용 기능을 사용하며, Android 단말은 배정 운행을 받아 시작·종료하고 영상과 센서를 전송할 수 있어야 한다.")
    w.h("2.4 제약사항")
    w.p("도로 데이터의 적용 지역은 부산이며, 차량 프로필은 시연을 위한 구성이다. Live View는 한 번에 하나의 Android 스트림을 다룬다. GPU 추론과 실제 단말 송출은 해당 장비에서 검증한다. 실제 차량의 운행 중 경로 재배정은 이번 범위에서 제외한다.")

    w.page("3. 시스템 설계",True)
    w.image("system_architecture.png","ITS 통합 시스템 구성")
    w.table("계층별 구성과 책임",["계층","구성요소","책임"],LAYERS,[2.8,5.7,8.9])
    w.p("서비스는 하나의 저장소에서 관리하되 런타임을 분리한다. Node는 업무 상태와 저장을 담당하고 Routing은 계산을 담당한다. Android의 영상은 Go 릴레이를 통해 Vision으로 전달하고, 녹화 객체는 MinIO에 저장한다.")

    w.page("3. 시스템 설계 — 데이터와 AI 연동")
    w.h("3.1 데이터 설계")
    w.image("domain_structure.png","실제 운행·가상 배차 데이터 분리",16.2)
    w.p("실제 GPS는 vehicle_position에, 가상 위치는 virtual_vehicle_state에 저장한다. 영상 세그먼트와 탐지 샘플은 원본 영상의 PTS, relay epoch와 frame sequence를 사용하여 연결한다. 같은 차량 테이블을 사용하더라도 출처와 운행 상태를 구분한다.")
    w.h("3.2 AI 모델 연동")
    w.table("주요 AI 모델 역할",["모델","역할","연동 결과"],[["Custom YOLO26-seg","객체 탐지·인스턴스 분할","클래스·박스·객체 마스크"],["UniDepth","영상 기반 거리 측정","객체 거리 정보"]],[5,5.6,6.8])
    w.p("분할 결과를 거리 정보와 연결하여 관제 화면에 제공한다. 객체 클래스별로 박스 또는 분할 표시를 선택할 수 있게 구성한다.")

    w.page("4. 개발 환경",True)
    w.table("서비스별 개발 기술",["영역","기술","용도"],STACK,[3,7,7.4])
    w.h("4.1 개발·관리 도구")
    w.p("Git으로 코드를 관리하고, Node는 npm, Python 서비스는 uv, 미디어 릴레이는 Go module, Android는 Gradle로 의존성을 관리한다. API 계약은 OpenAPI로 확인하고, 데이터베이스 변경은 Prisma 마이그레이션과 ERD 문서에 반영한다.")
    w.h("4.2 모델 개발 환경")
    w.table("분할 모델 학습 계획 기준",["항목","기준"],[["기반 모델","YOLO26s-seg"],["구조 실험","A0 기본 / A1 CARAFE / A2 ASPP / A3 CARAFE+ASPP / A4 디코더 결합"],["데이터셋","COCO 2017 인스턴스 분할"],["기본 입력 크기","640 × 640"],["비교 기준","Mask Precision·Recall·mAP50·mAP50-95"],["실험 기록","학습 설정, CSV, 가중치, 평가 그래프"]],[4,13.4])

    w.page("4. 개발 환경 — UniDepth 모델 작성란")
    w.h("4.3 거리 측정 모델")
    w.p("UniDepth를 거리 측정 모델로 사용한다. 상세 모델 자료는 아래 표에 추가한다.")
    w.table("UniDepth 모델 구성·학습 설정",["항목","내용"],BLANK_MODEL_ROWS,[4.8,12.6],blank=True)
    w.blank_area("UniDepth 모델 구조도·처리 흐름 작성 공간",5.0)

    w.page("5. 개발 계획",True)
    w.table("영역별 개발 계획",["단계","Frontend","Backend","AI"],[
        ["요구사항·설계","관제 화면·사용자 흐름 정리","서비스 책임·API·DB 설계","데이터셋·클래스·평가 기준 정리"],
        ["기반 구현","로그인·지도·차량 선택","인증·차량·운행 API, BIMS 연동","기본 분할 모델과 추론 연결"],
        ["영상·기록","Live View·분할 표시·재생 UI","WebRTC 릴레이·녹화·메타데이터","CARAFE·ASPP·디코더 실험, 거리 연동"],
        ["경로·배차","경로 미리보기·도로 통제 UI","차량 제약 A*, 가상 배차·시뮬레이션","교통 객체 성능 비교·추론 최적화"],
        ["검증·정리","시연 흐름·오류 상태 확인","통합·권한·기록 검증","성능 표·그래프·최종 문서 정리"],
    ],[2.2,4.7,5.7,4.8],font_size=8.8)
    w.h("5.1 일정 및 담당")
    w.table("일정·담당자 작성란",["작업","기간","담당자"],[["요구사항·설계","",""],["관제·경로·배차","",""],["영상·녹화·재생","",""],["분할 모델 개발·평가","",""],["거리 모델 정리·검증","",""],["통합 검증·시연","",""]],[7,5.2,5.2],blank=True)
    w.h("5.2 검증 계획")
    w.p("역할별 접근 제어, 경로 제한 반영, GPS 소스 전환, 영상·탐지 동기화, 가상 배차 수락과 차단 도로 대응을 확인한다. 분할 모델은 같은 평가 조건에서 비교하고, 거리 측정은 정답 거리 자료를 확보한 후 평가한다.")

    w.page("6. 시연 시나리오",True)
    w.h("6.1 차량 관제 및 영상 분석")
    w.steps(["운영자가 로그인한 뒤 지도에서 차량을 선택한다.","차량 위치와 계획 경로를 확인하고 Live View를 연다.","Android에서 배정 운행을 시작하고 카메라·GPS/IMU를 송출한다.","객체 박스·분할·거리 표시를 확인하고 클래스별 표시를 바꾼다.","GPS가 지연되는 상황에서 지연 상태와 선택적 기록 보정 동작을 확인한다."])
    w.h("6.2 가상 배차 및 도로 통제")
    w.steps(["가상 시나리오에서 차량·출발지·목적지·경유지를 선택한다.","경로를 미리 보고 모의 운전자에게 배차를 요청한다.","수락 후 차량이 이동하는 모습을 확인한다.","정체 또는 차단 구역을 그리고 차량별 자동 추종 설정에 따른 대응을 확인한다.","경유지를 변경하고 현재 위치에서 다시 계산한 경로를 확인한다."])
    w.h("6.3 운행 기록 재생")
    w.steps(["종료한 운행의 녹화 목록을 선택한다.","녹화 영상을 재생하고 해당 시점의 탐지 결과와 GPS를 확인한다.","타임라인을 이동하며 영상·지도·탐지 결과가 함께 갱신되는지 확인한다."])

    w.page("7. 결론",True)
    w.h("7.1 프로젝트 요약")
    w.p("본 프로젝트는 차량의 위치와 경로, 카메라 영상과 AI 분석, 운행 기록을 하나의 관제 흐름으로 연결한다. Node의 업무 데이터 관리와 Python의 경로·비전 계산, Go의 미디어 처리를 분리하여 각 서비스의 역할을 명확히 한다.")
    w.p("Custom YOLO26-seg는 CARAFE, ASPP와 DeepLabV3+ 방식의 디코더를 실험하여 객체 분할 성능을 비교한다. UniDepth는 거리 측정을 담당하며 상세 구성과 평가 결과는 해당 표에 보완한다. 가상 배차는 실제 운행과 분리된 환경에서 도로 통제와 우회 대응을 검토하는 기능으로 구성한다.")
    w.h("7.2 향후 개선 방향")
    for t in ["부산 실제 주행 영상과 작은 객체·가림·야간 장면을 확보하여 분할 성능을 보완한다.","카메라 보정과 정답 거리 자료를 정리하고 거리 측정의 오차를 평가한다.","지도·영상·탐지의 동기화와 장시간 송출·녹화 안정성을 검증한다.","도로 제한 데이터의 정확도·갱신 체계를 개선하고 실제 운행 적용 가능성을 검토한다."]:
        w.bullet(t)
    w.h("7.3 작성 근거")
    w.note("시스템 구성: P4 통합 백엔드 시스템 설계 문서 및 현재 저장소. 분할 모델 구성: 제공된 yolo_carafe_aspp 소스와 실험 설정. 이전 프로젝트 PDF는 문서 구성과 서식 참고에 사용하였다.")
    return w.save("4차_프로젝트_기획서.docx")


def report():
    w=Writer("결과 보고서");w.cover();w.toc()
    w.page("1. 프로젝트 개요",True)
    w.h("1.1 프로젝트명");w.p(PROJECT)
    w.h("1.2 프로젝트 목적")
    w.p("본 프로젝트는 차량의 현재 위치와 계획 경로, 카메라 영상의 객체 분석 결과, 과거 운행 기록을 하나의 웹 관제 화면에서 확인하는 것을 목적으로 한다. 차량 제원을 고려한 경로 계산과 가상 배차·도로 통제 시뮬레이션을 제공하여 관제 업무와 상황 대응의 검토를 지원한다.")
    w.h("1.3 개발 배경")
    w.p("차량 위치, 운행 경로와 도로 영상이 서로 분리되어 있으면 운영자가 현재 상황을 연결해 판단하기 어렵다. 대형 차량은 통과 가능한 도로의 제약도 고려해야 한다. 이에 지도·경로·영상·기록을 연동하고, 실제 운행에 영향을 주지 않는 가상 환경에서 배차와 도로 통제를 검토하도록 시스템을 구성하였다.")
    w.h("1.4 프로젝트 범위")
    w.table("프로젝트 기능 영역",["영역","구성 기능","주요 서비스"],FEATURES,[2.5,8.1,6.8])
    w.note("본 보고서의 구현 설명은 현재 코드·설정과 설계 문서를 기준으로 한다. 분할 성능은 보관된 학습·평가 기록을 사용한다.")

    w.page("1.5 요구사항 분석 및 제약사항")
    w.table("주요 요구사항과 구현 구성",["구분","요구사항","구현 구성"],[
        ["접근 제어","역할에 맞는 기능 사용","JWT 인증, ADMIN/OPERATOR/VIEWER 인가"],
        ["차량 관제","여러 차량의 위치·계획 경로 조회","BIMS·단말 GPS 병합, Leaflet 지도"],
        ["경로 탐색","차량 제약을 반영한 경로 계산","높이·중량·너비·접근·회전 제한과 A*"],
        ["영상 분석","실시간 영상과 객체 결과 조회","Android·Relay·Vision·Live View 연결"],
        ["기록 재생","영상·탐지·GPS 동기화 조회","PTS·epoch·sequence 기반 저장·재생"],
        ["가상 배차","배차·경유지·도로 통제 시연","시나리오별 가상 상태와 경로 버전 관리"],
        ["운영","서비스 일괄 실행과 보안 접점","Docker Compose, Nginx TLS, Coturn"],
    ],[2.5,6.2,8.7])
    w.h("1.6 사용자와 권한")
    w.table("사용자별 역할",["사용자","역할","관심사"],STAKEHOLDERS,[3,8,6.4])
    w.h("1.7 제약사항")
    w.p("지역 범위는 부산 도로망이며, 차량 프로필은 시연용이다. Live View는 단일 Android 스트림을 다룬다. GPU 추론·실제 Android 송출은 원래 장비에서 검증해야 한다. 실제 차량의 운행 중 경로 재배정, SUMO와 Tauri는 현재 통합 범위에서 제외된다.")

    w.page("2. 개발 환경",True)
    w.h("2.1 개발 언어 및 기술 스택")
    w.table("서비스별 개발 기술",["영역","기술","용도"],STACK,[3,7,7.4])
    w.h("2.2 주요 라이브러리")
    w.p("Node 백엔드는 Express 5, Prisma 7과 PostgreSQL 어댑터를 사용한다. Zod로 요청을 검증하고 jose·Argon2로 토큰과 비밀번호를 처리하며 Pino로 로그를 기록한다. 관제 웹은 별도 프레임워크 빌드 과정 없이 JavaScript와 Leaflet로 구성된다.")
    w.p("Vision은 Python 3.12, FastAPI, PyTorch, 커스텀 Ultralytics와 영상 처리 라이브러리를 사용한다. 현재 잠금 파일에는 PyTorch 2.12.1+cu130, torchvision 0.27.1+cu130이 기록되어 있다. 서비스의 패키지 설정은 모델 학습 환경과 구분하여 관리한다.")
    w.h("2.3 개발 도구")
    w.p("서비스별로 npm, uv, Go module과 Gradle을 사용한다. Git으로 변경을 관리하고 OpenAPI로 업무 API를 문서화한다. 테스트 구성에는 Vitest·Supertest, Python 테스트, Go 테스트와 Android 단위 테스트가 포함되어 있다.")
    w.note("근거: node/package.json, 각 서비스 pyproject.toml·uv.lock, Go module, Android Gradle 설정. 표의 버전은 저장소 설정 기준이다.")

    w.page("2.4 데이터 저장소 및 운영 환경")
    w.table("데이터 저장 구조",["저장소","저장 데이터","담당·용도"],[
        ["PostgreSQL/PostGIS","계정, 차량, 운행, 경로, 위치, 가상 상태","Node의 업무 데이터 및 공간 쿼리"],
        ["MinIO","녹화 영상 세그먼트 객체","Relay 업로드, 브라우저 재생"],
        ["Routing 메모리","도로 그래프, 위치 스냅샷, 계산 캐시","경로 계산과 일시적 위치 추적"],
        ["routing_state","선택한 BIMS 실시간·재생 소스와 보정 설정","소스 설정의 재시작 후 유지"],
        ["Relay 임시 저장","녹화·업로드 저널과 미디어 버퍼","영상 처리·업로드 상태 관리"],
    ],[3.3,7.5,6.6])
    w.h("2.5 운영 및 배포 환경")
    w.table("주요 서비스의 외부 접점",["서비스","접점","외부 노출"],[
        ["Nginx","39001 관제 / 39002 영상·단말 / 39003 저장소","HTTPS"],
        ["Coturn","39004–39007","TURN"],
        ["Node / Routing","3000 / 8000","내부 네트워크"],
        ["Vision / Relay","39011 / 39012","내부 네트워크"],
        ["PostgreSQL / MinIO","5432 / 9000","내부 네트워크"],
    ],[5.1,7.2,5.1])
    w.p("운영·개발 Compose 설정을 구분한다. 운영용 Vision 코드는 이미지에 포함되고 모델 디렉터리는 읽기 전용 디렉터리 바인드 마운트로 연결된다. Compose의 기본 분할 모델 경로는 a4_best.engine이며, Python 단독 실행 설정의 기본 경로는 yolo26s-seg.pt이다.")
    w.p("관제 AI 도우미는 별도 LLM 서비스와 외부 검색 저장소를 사용하는 부가 기능이다. 핵심 영상 모델 설명은 Custom YOLO26-seg와 UniDepth를 중심으로 정리한다.")

    w.page("3. 시스템 설계",True)
    w.h("3.1 전체 시스템 구성")
    w.image("system_architecture.png","통합 시스템 서비스 구성")
    w.table("계층별 구성과 책임",["계층","구성요소","책임"],LAYERS,[2.8,5.7,8.9])
    w.p("Node가 업무·영속 데이터를 관리하고, Routing은 그래프와 일시적 추적 상태를 관리한다. 미디어 처리와 AI 추론은 Node를 경유하는 고빈도 영상 전달 방식으로 합치지 않고 별도 서비스에서 수행한다.")
    w.note("추가 제공된 시스템 설계 문서의 서비스 경계를 적용하되, 전달 프로토콜과 부가 서비스는 현재 코드·Compose 설정에 맞춰 정리하였다.")

    w.page("3.2 서브시스템과 통신 인터페이스")
    w.table("주요 서브시스템",["서브시스템","핵심 책임","인터페이스"],[
        ["Node/Express","인증·차량·운행·기록·배차, API 파사드","/api/v1/*, /internal/*"],
        ["Routing/Tracking","BIMS, 단말 위치 병합, A*, 통제 간선 계산","/internal/vehicles, /internal/routing/*"],
        ["Media Relay","WebRTC 수신, 세션 검증, 녹화·GPS 전달","/offer/android, Unix 소켓, S3"],
        ["Vision","객체 분할·추적·거리, 프레임·센서 매칭","추론 feed, /ws/playback, 내부 API"],
        ["관제 웹","지도·차량·경로·영상·재생·가상 배차","HTTPS REST, iframe, postMessage"],
        ["Android","배정 운행 제어, 카메라·GPS/IMU 송출","HTTPS REST, WebRTC DataChannel"],
    ],[3.2,7.6,6.6])
    w.h("3.3 주요 데이터 흐름")
    w.table("서비스 간 데이터 전달",["출발 → 도착","전달 방식","내용"],[
        ["관제 웹 → Node","HTTPS REST + JWT","인증·조회·운행·배차 요청"],
        ["Node → Routing","내부 HTTP","위치 스냅샷·경로·통제 계산"],
        ["Android → Relay","WebRTC / DataChannel","H.264, GPS/IMU, QR 이벤트"],
        ["Relay → Vision","Unix 소켓 / 내부 HTTP","H.264 feed와 프레임 제어 / 센서"],
        ["Relay → Node / MinIO","내부 HTTP / S3","세션 검증·GPS·세그먼트 / 영상"],
        ["Vision → Node","내부 HTTP","녹화 탐지 샘플"],
        ["브라우저 ↔ Vision / MinIO","WebSocket / presigned GET","Live View 데이터 / 녹화 영상"],
    ],[4.5,5,7.9],font_size=8.8)
    w.p("현재 Relay의 yolofeed는 Unix 소켓으로 Vision과 연결된다. 설계 문서의 미디어·비전 서비스 역할을 유지하면서 실제 전달 경로를 구조도와 인터페이스 표에 반영하였다.")

    w.page("3.4 데이터베이스 설계")
    w.image("domain_structure.png","실제 운행과 가상 배차의 주요 관계",16.5)
    w.table("핵심 데이터 그룹",["그룹","주요 테이블","내용"],[
        ["계정·차량","platform_account, vehicle, driver","사용자 역할·차량 출처·제원"],
        ["운행·경로","trip, route, vehicle_position","운행 상태·버전 경로·GPS"],
        ["영상 기록","trip_video, trip_video_detection_sample","영상 객체·PTS·탐지 JSON"],
        ["가상 운영","virtual_scenario, virtual_dispatch_request","시나리오·배차 요청·수락 상태"],
        ["가상 주행","virtual_trip, virtual_route, virtual_vehicle_state","경로 버전·현재 간선·이동 상태"],
        ["가상 통제","virtual_trip_waypoint, virtual_road_restriction, virtual_operator_event","경유지·통제·감사 기록"],
    ],[2.5,8.9,6])
    w.p("실제 위치와 가상 위치는 별도 테이블로 관리한다. 공간 데이터는 PostGIS geography를 사용하고, 관련 읽기·쓰기는 TypedSQL 등 공간 SQL로 처리한다. 상태값은 문자열과 데이터베이스 CHECK 제약으로 관리한다.")

    w.page("3.5 데이터 정합성 및 화면 설계")
    w.h("3.5.1 데이터 정합성")
    w.table("주요 식별·정합성 기준",["항목","기준","목적"],[
        ["차량·경로·위치 출처","vehicle_source / route_source / telemetry_source","출처 혼동 방지"],
        ["단말 GPS","recording_session_id + source_timestamp_ns","같은 GPS의 멱등 저장"],
        ["영상·탐지","recordingSession + relay epoch + frame sequence + PTS","원본 영상과 결과 연결"],
        ["가상 배차","idempotency_key와 트랜잭션","중복 수락·운행 생성 방지"],
        ["가상 경로·명령","routeVersion / commandVersion / sequence","오래된 갱신 적용 방지"],
    ],[3.1,8.9,5.4])
    w.p("object_class, detection_event, event_image, route_deviation, alert와 transport_goal은 확장을 위한 스키마에 포함되어 있다. 현재 녹화 탐지 결과의 저장 경로는 trip_video_detection_sample이므로, 스키마 정의와 운영 기능을 구분한다.")
    w.h("3.5.2 관제 화면")
    w.table("워크스페이스별 주요 화면",["화면","표시·조작 내용"],[
        ["일반 관제","차량 지도, 선택 차량, 운행·계획 경로, Live View"],
        ["운행 기록","녹화 목록, 영상 타임라인, 탐지·GPS 재생"],
        ["가상 배차","차량·출발지·도착지·경유지, 경로 미리보기, 배차 요청"],
        ["도로 통제","정체·차단 구역, 영향 도로, 차량별 자동 추종"],
        ["관제 AI 도우미","현재 차량 상태를 포함한 질의와 근거 기반 답변"],
    ],[4.2,13.2])
    w.p("GPS 지연 상태와 데이터 소스는 화면에 구분하여 표시한다. 기록 경로를 이용한 위치 보정은 운영자가 활성화한 경우에 적용한다. 화면에 표시되는 추정·재생 위치를 실제 수신 GPS와 혼동하지 않도록 출처를 유지한다.")

    w.page("3.6 영상·센서 동기화와 가상 주행")
    w.h("3.6.1 영상·GPS/IMU 처리")
    w.steps(["Android가 배정 운행과 녹화 세션으로 WebRTC 연결을 요청한다.","Relay가 Node에서 세션의 trip·vehicle을 검증하고 영상과 센서를 수신한다.","GPS는 Node에 저장하며 GPS/IMU는 Vision으로 전달한다.","Vision이 원본 프레임 시각과 센서를 매칭하고 탐지 결과를 생성한다.","녹화 영상과 탐지 샘플을 같은 PTS·epoch·sequence 기준으로 재생한다."])
    w.h("3.6.2 가상 배차·도로 통제")
    w.table("가상 운영 주요 동작",["상황","처리"],[
        ["배차 수락","요청 수락 후 가상 운행·경로·상태를 생성하고 주행 시작"],
        ["자동 추종 ON","통제 변화 시 남은 경로를 다시 계산해 적용"],
        ["자동 추종 OFF","정체에서는 기존 주행 유지, 차단 앞에서는 운영자 대응 대기"],
        ["OFF → ON","현재 위치와 남은 경유지·목적지를 기준으로 즉시 재계산"],
        ["점유 도로 차단 요청","가상 차량이 점유한 도로를 포함하면 409 ROAD_OCCUPIED"],
        ["경로 없음","NO_ROUTE 등 상태로 표시하고 운영자 대응 지원"],
    ],[4.5,12.9])
    w.p("가상 시뮬레이션은 브라우저 연결과 분리된 서버 상태로 진행한다. 실제 운행 영상·녹화와 가상 차량의 이동을 함께 사용하는 데이터 경로로 합치지 않는다.")

    w.page("4. AI 모델 및 구현",True)
    w.h("4.1 주요 AI 모델")
    w.table("주요 모델과 역할",["모델","입력","역할·결과"],[["Custom YOLO26-seg","영상 프레임","객체 클래스·박스·인스턴스 마스크"],["UniDepth","영상 프레임","거리 측정"]],[5,4,8.4])
    w.h("4.2 Custom YOLO26-seg 실험 구성")
    w.p("분할 모델은 YOLO26s-seg를 기준으로 업샘플링, 문맥 특징과 마스크 디코더를 변경하였다. 구조별 구현을 A0~A4로 구분하고, 저장된 학습 기록을 사용하여 성능을 비교한다.")
    w.table("분할 모델 구조별 구성",["실험","구조","변경 내용"],[
        ["A0","YOLO26s-seg 기본","기본 백본·Neck·분할 헤드"],
        ["A1","CARAFE","Top-down 업샘플링에 내용 기반 특징 재조합"],
        ["A2","ASPP","고수준 특징에 병렬 팽창 합성곱 문맥 결합"],
        ["A3","CARAFE + ASPP","내용 기반 업샘플링과 다중 문맥 결합"],
        ["A4","CARAFE + ASPP + DeepLabV3+ 방식 디코더","P2의 상세 특징과 고수준 특징을 결합하여 Proto 생성"],
    ],[1.5,6.2,9.7])
    w.h("4.3 모델 변경 목적")
    w.p("CARAFE는 업샘플링 과정에서 입력 특징의 내용을 반영한다. ASPP는 서로 다른 수용 영역의 문맥을 결합한다. A4의 디코더는 낮은 stride의 상세 특징을 마스크 생성에 사용하여 객체 경계를 보완하도록 구성하였다. 각 변경의 실제 효과는 성능 평가에서 확인한다.")

    w.page("4.4 A4 분할 모델 구조")
    w.image("custom_yolo_architecture.png","CARAFE·ASPP·DeepLabV3+ 방식 디코더를 결합한 A4")
    w.table("A4 주요 모듈 설정",["모듈","설정","역할"],[
        ["ASPP","P5 / stride 32, dilation 3·6·9","병렬 1×1·팽창 3×3·전역 풀링 문맥 결합"],
        ["Neck CARAFE","scale 2, 재조합 kernel 5, encoder kernel 3","P5→P4와 P4→P3 업샘플링"],
        ["Low-level projection","P2 / stride 4, 1×1 conv, 48채널","상세 특징의 채널 축소"],
        ["Decoder","P3 특징 CARAFE 업샘플링, concat, 3×3 conv 두 번","P2와 문맥 특징 정제"],
        ["Proto 출력","32개 마스크 prototype","객체별 계수와 결합할 분할 표현"],
        ["검출 경로","P3·P4·P5의 기존 검출 구조 유지","클래스·박스·마스크 계수 생성"],
    ],[3.4,6.5,7.5])
    w.p("학습 시 보조 semantic 분기가 사용된다. 최종 객체별 마스크는 Proto와 객체별 마스크 계수로 구성되며, 디코더의 결합 대상은 P2 상세 특징과 Neck에서 얻은 P3 특징이다.")
    w.note("근거: experiments/configs/yolo26-seg-carafe-aspp-deeplabv3plus.yaml 및 ultralytics/nn/modules/custom_seg.py.")

    w.page("4.5 탐지 클래스와 학습 설정")
    w.p("학습은 COCO 2017의 80개 클래스를 기준으로 수행하였다. 관제에서 주로 확인하는 교통 관련 클래스는 다음 여섯 종류이며, 화면의 클래스 선택은 체크포인트의 클래스 이름을 기준으로 처리한다.")
    w.table("주요 교통 객체 클래스",["COCO ID","클래스","표시명"],CLASSES,[3,7.2,7.2])
    w.table("A0~A4 공통 학습 설정",["항목","설정값"],[
        ["데이터 설정","coco.yaml / COCO 2017"],["입력 크기","640 × 640"],
        ["학습 epoch / patience","100 / 100"],["배치 크기","32"],
        ["Optimizer 설정","Ultralytics auto"],["학습률 설정","lr0 = 0.01, lrf = 0.01"],
        ["Seed / 데이터 비율","0 / 1.0"],["정밀도","AMP 활성화"],
        ["학습 방식","실험별 단일 GPU 실행"],["기록","args.yaml, results.csv, 가중치 및 그래프"],
    ],[5,12.4])
    w.note("설정은 각 실행 폴더의 args.yaml을 기준으로 정리하였다. 학습 장비 설명과 서비스 실행 의존성은 별도 환경으로 구분한다.")

    w.page("4.6 실시간 추론 및 관제 연동")
    w.p("Vision 서비스는 분할 모델인지 확인한 뒤 추론을 시작한다. 클래스 설정은 체크포인트의 클래스 맵으로 해석하며, 추적 모드에서는 ByteTrack 설정과 유지 상태를 사용한다. 프레임의 박스·클래스·confidence·마스크와 추적 ID를 관제용 결과로 정리한다.")
    w.table("추론·표시 설정",["항목","현재 구성"],[
        ["모델 형식","분할 체크포인트 또는 TensorRT engine"],
        ["Compose 기본 가중치","a4_best.engine"],
        ["Python 설정 기본 가중치","yolo26s-seg.pt"],
        ["입력 크기","auto 또는 지정 크기, 최대 크기 설정"],
        ["정밀도","GPU half 설정 또는 모델에 맞는 정밀도"],
        ["결과","class, confidence, bbox, mask, track_id, distance_m, 상태"],
        ["표시 방식","클래스별 박스 / 분할 선택"],
        ["녹화 결과 저장","Node 내부 API → trip_video_detection_sample"],
    ],[5,12.4])
    w.h("4.6.1 동기화 및 실패 상태")
    w.p("결과에는 원본 프레임 식별 정보가 포함되며 프레임과 탐지의 연결은 도착 순서만으로 처리하지 않는다. 거리 정보가 없는 경우 숫자를 임의로 채우지 않고 상태와 함께 전달한다. 녹화 저장 경로는 추론과 분리된 큐로 처리한다.")
    w.h("4.6.2 관제 AI 도우미")
    w.p("부가 기능으로 현재 차량 스냅샷과 운송 관련 문서 근거를 활용하는 관제 AI 도우미가 구성되어 있다. Node가 별도 LLM 서비스에 요청하고, 검색 문맥을 활용한 한국어 답변을 관제 패널에 표시한다.")
    w.note("서비스 구성은 코드·설정 확인 결과이며, 본 문서 작성 중 실제 GPU 추론이나 Android 송출을 수행한 결과는 아니다.")

    w.page("4.7 UniDepth 거리 측정 모델")
    w.p("UniDepth는 영상 기반 거리 측정을 담당한다. 상세 모델 구성과 학습·실행 자료를 추가할 수 있도록 아래 표를 마련하였다.")
    w.table("UniDepth 모델 구성·학습 설정",["항목","내용"],BLANK_MODEL_ROWS,[4.8,12.6],blank=True)
    w.blank_area("UniDepth 모델 구조도·연동 흐름 작성 공간",5)

    w.page("5. 성능 평가",True)
    w.h("5.1 평가 지표와 기록 기준")
    w.table("분할·탐지 평가 지표",["지표","의미"],[
        ["Precision","모델의 예측 중 정답으로 판단된 비율"],
        ["Recall","정답 객체 중 모델이 찾아낸 비율"],
        ["Box mAP50","박스 IoU 0.50 기준의 클래스별 AP 평균"],
        ["Box mAP50-95","박스 IoU 0.50~0.95 기준의 AP 평균"],
        ["Mask mAP50","객체 마스크 IoU 0.50 기준의 AP 평균"],
        ["Mask mAP50-95","마스크 IoU 0.50~0.95 기준의 AP 평균"],
        ["분할 손실","학습·검증에서 계산한 분할 오차 항목"],
    ],[5,12.4])
    w.h("5.2 평가 자료 구분")
    w.p("첫 번째 비교는 저장된 A0~A4 학습 검증 기록이다. 실행별 Mask mAP50-95가 가장 높은 epoch를 선택하고, 해당 행의 모든 지표를 함께 사용한다. 표의 최고 epoch는 이 선택 기준을 뜻하며, 가중치 파일의 별도 선정 기준과 구분한다.")
    w.p("두 번째 비교는 기존 모델 비교 문서에 기록된 공통 조건 검증이다. A0, A4와 공식 체크포인트를 같은 COCO val2017 조건에서 평가한 자료를 정리하였다. 학습 검증 기록과 공통 조건 결과는 서로 다른 평가 기록이다.")
    w.note("데이터 출처: 각 실행 폴더의 results.csv 및 experiments/notes/yolo26_seg_model_comparison.md. 새 학습·검증은 실행하지 않았다.")

    w.page("5.3 A0~A4 학습 기록 비교")
    w.table("실행별 최고 epoch의 Mask 지표",["실험","Epoch","Precision","Recall","mAP50","mAP50-95"],
            [[f"A{i}",str(int(b['epoch']))]+[f"{b[k]:.5f}" for k in METRICS[4:]] for i,b in enumerate(BEST)],
            [1.6,1.8,3.5,3.5,3.5,3.5],font_size=9)
    w.table("동일 epoch의 Box 지표",["실험","Epoch","Precision","Recall","mAP50","mAP50-95"],
            [[f"A{i}",str(int(b['epoch']))]+[f"{b[k]:.5f}" for k in METRICS[:4]] for i,b in enumerate(BEST)],
            [1.6,1.8,3.5,3.5,3.5,3.5],font_size=9)
    w.image("best_comparison.png","실험별 최고 Mask mAP50-95",16.5)
    w.p("저장된 학습 기록에서 A4의 Mask mAP50-95는 0.39683으로 가장 높다. A0의 0.39474와 비교하면 0.00209, 즉 약 0.209%p 높다. A1의 Mask Recall은 0.56126으로 다섯 실행 중 가장 높다.")
    w.note("비교 기준: 각 실행의 최고 Mask epoch 행. 수치는 원본 CSV 값을 소수점 다섯 자리로 표시하였다.")

    w.page("5.4 분할 성능 변화")
    w.image("training_comparison.png","A0~A4의 epoch별 Mask mAP50-95",17)
    w.table("최고 epoch와 마지막 epoch 비교",["실험","최고 Epoch","최고 Mask mAP50-95","100 Epoch Mask mAP50-95"],
            [[f"A{i}",str(int(b['epoch'])),f"{b['metrics/mAP50-95(M)']:.5f}",f"{r[-1]['metrics/mAP50-95(M)']:.5f}"]
             for i,(b,r) in enumerate(zip(BEST,ALL_ROWS))],[2,3,6.2,6.2])
    w.p("각 실행은 100 epoch의 기록을 갖고 있으며 최고 Mask 성능은 91~92 epoch에 나타난다. 마지막 epoch 수치와 최고 수치를 함께 제시하여 학습 종료 시점과 성능 선택 시점을 구분하였다.")
    w.p("A4는 최고 epoch 이후 학습 분할 손실이 감소하지만 검증 분할 손실은 증가하고 Mask mAP50-95는 낮아진다. 후속 실험에서는 단순 학습 연장보다 검증 성능을 기준으로 한 체크포인트 선정과 제한된 미세 조정을 검토할 수 있다.")
    w.note("그래프는 저장된 CSV를 사용해 다시 작성한 것이며, 보간·평활화 없이 epoch별 값을 표시하였다.")

    for i,(rows,best) in enumerate(zip(ALL_ROWS,BEST)):
        w.page(f"5.5.{i+1} A{i} 학습 결과")
        w.p(["A0는 YOLO26s-seg 기본 구조의 비교 기준이다.","A1은 CARAFE를 적용하여 업샘플링 단계의 특징 재조합을 변경한 구조이다.","A2는 ASPP를 적용하여 고수준 특징에 다중 문맥을 결합하는 구조이다.","A3는 CARAFE와 ASPP를 함께 적용한 구조이다.","A4는 CARAFE·ASPP와 DeepLabV3+ 방식의 마스크 디코더를 결합한 구조이다."][i])
        w.image(f"training_A{i}.png",f"A{i} 분할 정확도·mAP·손실·박스 성능 기록",17)
        w.table(f"A{i} 최고 Mask epoch의 성능",["구분","Precision","Recall","mAP50","mAP50-95"],
                [["Box"]+[f"{best[k]:.5f}" for k in METRICS[:4]],
                 ["Mask"]+[f"{best[k]:.5f}" for k in METRICS[4:]]],[2.8,3.65,3.65,3.65,3.65])
        w.p(f"저장된 학습 기록은 100 epoch이며, Mask mAP50-95 최고 시점은 {int(best['epoch'])} epoch이다. 해당 값은 {best['metrics/mAP50-95(M)']:.5f}이고, 100 epoch 값은 {rows[-1]['metrics/mAP50-95(M)']:.5f}이다.")
        w.note(f"자료: {RUN_NAMES[i]}/results.csv. 점선은 최고 Mask mAP50-95 epoch를 나타낸다.")

    w.page("5.6 공통 조건의 모델 평가")
    w.table("공통 조건 평가 설정",["항목","설정"],[
        ["비교 모델","A0, A4, 공식 yolo26s-seg.pt"],
        ["검증 데이터","COCO val2017 / 5,000 이미지"],
        ["입력 크기 / 배치","640 / 4"],
        ["장비 / 정밀도","Quadro P2200, GPU 0 / FP32"],
        ["Confidence / IoU / max detections","0.001 / 0.7 / 300"],
        ["NMS 설정","외부 NMS, nms=None"],
    ],[6,11.4])
    common = [
        ["A0","0.717","0.557","0.609","0.389"],
        ["A4","0.725","0.558","0.612","0.392"],
        ["공식 모델","0.716","0.564","0.618","0.393"],
    ]
    w.table("공통 조건의 전체 Mask 성능",["모델","Precision","Recall","mAP50","mAP50-95"],common,[3.4,3.5,3.5,3.5,3.5])
    w.table("공통 조건의 전체 Box 성능",["모델","Precision","Recall","mAP50","mAP50-95"],[
        ["A0","0.717","0.577","0.637","0.467"],["A4","0.728","0.580","0.640","0.470"],
        ["공식 모델","0.720","0.590","0.649","0.476"],
    ],[3.4,3.5,3.5,3.5,3.5])
    w.p("기존 평가 자료의 원정밀도 값 기준으로 A4의 전체 Mask AP50-95는 0.39161이며 A0 대비 0.00216 높다. 공식 체크포인트의 0.39285보다는 0.00124 낮다. 따라서 A4가 기본 학습 모델보다 개선된 부분과 공식 모델 대비 남은 차이를 함께 확인할 수 있다.")
    w.note("출처: 기존 yolo26_seg_model_comparison.md 및 A4_class_finetuning_plan.md. 표는 기존 공통 조건 결과를 소수점 세 자리로 표시하였다.")

    w.page("5.7 교통 객체 클래스별 성능")
    w.table("주요 클래스의 Mask AP50-95",["클래스","정답 객체 수","A0","A4","공식 모델","A4 − A0"],[
        ["person","10,777","0.466","0.465","0.468","−0.00137"],
        ["bicycle","314","0.201","0.196","0.212","−0.00539"],
        ["car","1,918","0.375","0.379","0.382","+0.00383"],
        ["motorcycle","367","0.368","0.378","0.388","+0.00910"],
        ["bus","283","0.650","0.649","0.654","−0.00095"],
        ["truck","414","0.348","0.356","0.356","+0.00776"],
    ],[3.6,3.0,2.5,2.5,2.8,3],font_size=8.8)
    w.image("class_comparison.png","주요 교통 객체의 클래스별 분할 성능",17)
    w.p("A4는 A0 대비 car·motorcycle·truck의 Mask AP50-95가 높고, person·bicycle·bus는 낮다. 특히 motorcycle의 증가 폭은 약 0.00910이다. truck은 공식 모델과 표의 반올림 값이 같으며, 원정밀도에서는 A4 0.35610과 공식 모델 0.35596으로 차이가 작다.")
    w.note("표의 AP는 소수점 세 자리, 차이는 기존 평가 자료의 원정밀도 기준이다. 정답 객체 수와 클래스별 지표는 기존 모델 비교 문서를 사용하였다.")

    w.page("5.8 UniDepth 성능 평가 작성란")
    w.p("거리 측정 평가 조건과 결과를 추가할 수 있도록 아래 표와 그래프 영역을 마련하였다.")
    w.table("UniDepth 평가 조건·성능",["항목","설정 / 측정값","비고"],BLANK_EVAL_ROWS,[5.4,7,5],blank=True)
    w.blank_area("거리 측정 결과·오차 그래프 작성 공간",6)

    w.page("5.9 시스템 검증 항목")
    w.table("통합 기능 검증 시나리오",["영역","확인할 시나리오","현재 근거"],[
        ["인증·권한","조회자 쓰기 요청 차단, 역할별 차량 API","라우터·미들웨어·테스트 코드"],
        ["차량 위치","BIMS/재생 전환, 지연 표시, 기록 보정","추적·지도·소스 설정 코드"],
        ["경로 계산","높이·중량·너비·회전 제한, 경로 없음","그래프·경로 코드와 테스트"],
        ["영상·센서","원본 시각 매칭, 탐지·GPS 동기화","타임라인·센서 테스트 코드"],
        ["녹화·재생","세그먼트 등록, PTS 조회, URL 발급","녹화 API·재생 테스트 코드"],
        ["가상 배차","중복 수락, 경유지 진행, 자동 추종 전환","배차·경유지·점유 테스트 코드"],
        ["도로 통제","점유 도로 차단 거부, 대체 경로·정지","가상 통제 구현과 테스트"],
    ],[2.6,8.9,5.9])
    w.p("저장소에는 위 동작을 확인하기 위한 테스트 코드가 포함되어 있다. 본 보고서 작성에서는 테스트 전체 실행이나 운영 장비의 통합 시연을 수행하지 않았으므로, 코드 확인을 테스트 통과 결과로 기재하지 않는다.")
    w.table("실행 환경별 성능 측정 작성란",["항목","장비 / 조건","측정값"],[["영상 송출 FPS","",""],["분할 추론 시간","",""],["통합 추론·표시 지연","",""],["경로 계산 시간","",""],["녹화·재생 동기화 오차","",""]],[6,6,5.4],blank=True)

    w.page("6. 프로젝트 결과",True)
    w.h("6.1 구현 결과")
    w.table("기능별 구현 구성",["영역","구현 내용","관련 구성"],[
        ["업무 API","인증, 차량 CRUD, 운행 생성·취소·단말 시작·종료","Node 모듈·PostgreSQL"],
        ["통합 지도","BIMS·단말 GPS, 계획 경로, 지연·소스 표시","Routing·관제 웹"],
        ["경로 탐색","OSM 그래프와 제한 데이터, 차량 제약 A*","Routing/Tracking"],
        ["영상 분석","WebRTC·추론 feed, 객체 분할·추적·거리 상태","Android·Relay·Vision"],
        ["운행 기록","세그먼트·탐지 저장, 재생 URL·타임라인","MinIO·Node·Vision·웹"],
        ["가상 배차","경로 미리보기, 요청 수락, 서버 시뮬레이션","Node·Routing·웹"],
        ["도로 통제","정체·차단, 점유 검사, 차량별 재경로","가상 상태·그래프 오버레이"],
        ["부가 도우미","차량 스냅샷과 문서 검색 기반 답변","Node·LLM 서비스"],
    ],[2.8,8.9,5.7])
    w.h("6.2 분할 모델 개발 결과")
    w.p("YOLO26s-seg를 기준으로 A0~A4 구조를 구성하고 학습 기록과 평가 그래프를 확보하였다. A4는 고수준 문맥과 저수준 상세 특징을 결합하는 분할 경로를 포함한다. 저장된 학습 검증 기록에서 A4는 전체 Mask mAP50-95 0.39683을 기록하였다.")
    w.p("공통 조건 평가에서는 기본 학습 모델 대비 개선이 확인되었으며, 클래스별 결과를 통해 성능이 개선된 객체와 후속 보완이 필요한 객체를 구분하였다. UniDepth의 상세 구성과 평가 결과는 본 문서의 빈 표에 추가하도록 마련하였다.")

    w.page("6.3 주요 기능 시연 — 관제·영상·기록")
    w.h("6.3.1 차량 관제와 실시간 영상")
    w.steps(["운영자가 로그인하고 지도에서 차량을 선택한다.","선택 차량의 현재 위치·상태·계획 경로를 확인한다.","Android가 배정 운행을 시작하고 WebRTC로 카메라와 센서를 보낸다.","Live View에서 객체 클래스·박스·마스크·거리 표시를 확인한다.","클래스별 표시 방식을 조절하고 해당 프레임의 GPS와 지도 위치를 확인한다."])
    w.h("6.3.2 GPS 지연과 소스 전환")
    w.steps(["실시간 BIMS 위치에서 마지막 관측 시각과 지연 상태를 확인한다.","기록 경로 보정 설정을 바꾸어 선택적 보정 동작을 확인한다.","BIMS 실시간과 재생 소스를 전환하고 지도·차량 상태 갱신을 확인한다."])
    w.h("6.3.3 운행 기록 재생")
    w.steps(["운행의 녹화 목록을 열고 재생할 영상 구간을 선택한다.","발급된 접근 URL로 영상을 불러온다.","영상 PTS를 기준으로 탐지 결과와 GPS를 표시한다.","타임라인을 이동하고 영상·탐지·지도 상태가 같은 구간을 가리키는지 확인한다."])
    w.note("시연 절차는 구현된 사용자 흐름을 정리한 것이다. 운영 장비에서 실행한 시연 결과는 별도로 기록한다.")

    w.page("6.4 주요 기능 시연 — 가상 배차·도로 통제")
    w.h("6.4.1 가상 배차")
    w.steps(["시나리오와 사용 가능한 가상 차량을 선택한다.","출발지·목적지·경유지를 입력하고 경로를 미리 확인한다.","모의 운전자에게 배차를 요청한다.","수동 또는 설정된 자동 수락 후 차량 이동을 확인한다.","가상 운행 상태와 경로 버전, 운영자 이벤트 기록을 확인한다."])
    w.h("6.4.2 정체·차단 대응")
    w.steps(["지도에서 정체 구역을 만들고 영향 도로를 확인한다.","자동 추종을 켠 차량이 새 경로를 계산하는 흐름을 확인한다.","차단 구역을 설정하고 자동 추종을 끈 차량의 대기 동작을 확인한다.","자동 추종 OFF→ON 전환 후 현재 위치에서 남은 경로를 다시 계산한다.","가상 차량이 점유한 도로를 차단하는 요청이 거부되는지 확인한다."])
    w.h("6.4.3 경유지 변경")
    w.p("운행 중 경유지를 추가·변경하고 현재 위치에서 목적지까지의 남은 경로를 확인한다. 실제 차량·녹화 기능과 가상 주행이 분리되어 동작하는지 함께 검토한다.")

    w.page("6.5 프로젝트 성과 및 한계점")
    w.h("6.5.1 프로젝트 성과")
    for t in ["업무·경로 계산·미디어·비전 서비스를 하나의 제품 흐름으로 연결하였다.","차량 출처와 위치 출처를 구분하고 실제·가상 운행 데이터를 분리하였다.","녹화 영상과 탐지·GPS 정보를 동일한 원본 시간축으로 조회하는 구조를 마련하였다.","가상 배차와 통제 상황을 서버 상태로 관리하여 관제 시나리오를 구성하였다.","분할 모델의 구조별 학습 기록과 교통 객체별 평가 자료를 정리하였다."]:
        w.bullet(t)
    w.h("6.5.2 한계점")
    w.table("주요 한계와 보완 항목",["항목","현재 한계","보완 방향"],[
        ["실환경 일반화","COCO 평가만으로 부산 주행 성능을 확정할 수 없음","주행·야간·가림·원거리 자료 확보"],
        ["객체별 성능","A4의 person·bicycle 등 성능 보완 필요","목표 클래스 미세 조정과 회귀 검증"],
        ["거리 측정","상세 모델·정답 거리 평가 자료 미기입","모델 자료와 거리 오차 결과 추가"],
        ["영상 확장","단일 Android 스트림 관제 구성","필요 시 스트림·세션 라우팅 확장"],
        ["도로 제약","시연용 차량 프로필·지역 데이터","현장 제한값 검증·데이터 갱신"],
        ["운영 검증","본 문서 작성에서 GPU·단말 E2E 미실행","원래 장비에서 통합·장시간 검증"],
    ],[3.1,7.6,6.7])

    w.page("7. 결론 및 향후 계획",True)
    w.h("7.1 프로젝트 요약")
    w.p("본 프로젝트는 차량 위치와 경로, 카메라 영상의 AI 분석과 운행 기록을 연결하는 통합 관제 시스템을 구성하였다. Node는 업무 데이터를 관리하고, Routing은 계산과 위치 추적을 담당하며, Go 릴레이와 Vision은 영상·추론 경로를 담당한다. 가상 배차·도로 통제는 실제 운행과 분리된 시뮬레이션으로 제공한다.")
    w.h("7.2 결론")
    w.p("Custom YOLO26-seg는 CARAFE·ASPP·디코더의 구조별 실험을 통해 분할 성능을 비교할 수 있는 자료를 확보하였다. A4의 성능은 기본 학습 모델보다 높아진 항목이 있지만, 전체 및 클래스별 결과를 함께 보아야 한다. UniDepth는 거리 측정을 담당하며 모델 상세와 평가 자료를 보완할 공간을 마련하였다.")
    w.h("7.3 향후 개선 방향")
    w.table("후속 개선 계획",["영역","개선 내용","확인 기준"],[
        ["분할 모델","실제 주행 데이터, person·bicycle 보완, 미세 조정","목표 클래스 AP와 전체 성능 회귀"],
        ["거리 모델","상세 모델 자료·카메라 보정·정답 거리 정리","거리대·장면별 오차"],
        ["영상·기록","장시간 송출·재연결·녹화 안정성","프레임 지연·누락·동기화 오차"],
        ["경로·통제","제한 데이터 검증과 계산 최적화","통행 제한 반영·경로 응답 시간"],
        ["업무 확장","경로 이탈·알림 등 스키마 기능 연결","실제 이벤트·권한·저장 검증"],
    ],[3.1,8.9,5.4])
    w.h("7.4 향후 활용 방안")
    w.p("차량 운영 교육과 관제 시연, 도로 통제 시나리오 검토에 활용할 수 있다. 실제 운행 적용을 위해서는 도로 제한과 거리 정확도, 운영 장비의 송출·녹화·재생 안정성을 추가로 검증해야 한다.")

    w.page("작성 근거 및 자료 위치")
    w.table("본문 작성에 사용한 자료",["자료","활용 범위"],[
        ["P4 통합 백엔드 시스템 설계 문서.md","요구사항·유스케이스·서비스 책임·DB 설계"],
        ["현재 저장소의 서비스 코드·Compose·Prisma 스키마","실제 구현·인터페이스·설정 확인"],
        ["yolo_carafe_aspp / experiments/configs 및 custom_seg.py","분할 모델 구조·모듈 설정"],
        ["yolo_carafe_aspp / runs/segment/experiments/results","A0~A4 args.yaml·results.csv·학습 자료"],
        ["experiments/notes/yolo26_seg_model_comparison.md","공통 조건 전체·클래스별 평가 결과"],
        ["experiments/notes/A4_class_finetuning_plan.md","원정밀도 클래스별 차이·후속 실험 방향"],
        ["이전 프로젝트 기획서·보고서 PDF","목차·서식·문체 참고"],
    ],[9,8.4],font_size=8.7)
    w.h("자료 수정 안내")
    w.p("본문과 표는 Word에서 수정할 수 있다. 학습 그래프는 source/data의 실행별 CSV에서 생성되며, 그림 원본은 source/assets에 있다. 수치를 변경할 때 성능 표·본문 설명·그래프를 함께 갱신한다. UniDepth 표와 그림 영역은 해당 자료를 추가하여 완성한다.")
    w.h("검증 범위")
    w.p("프로젝트 설명은 소스·설정·설계 자료를 확인한 결과이며, 모델 수치는 기존 기록을 정리한 것이다. 문서 작성 과정에서 새 학습, GPU 추론, 실차 송출 또는 운영 배포를 실행하지 않았다.")
    return w.save("4차_프로젝트_결과_보고서.docx")


if __name__ == "__main__":
    prepare_metrics();plot_training();plot_classes();diagram_arch();diagram_model();diagram_domain()
    manifest={"project":PROJECT,"documents":[planning(),report()]}
    (Path(__file__).parent/"document_manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print(json.dumps(manifest,ensure_ascii=False,indent=2))
